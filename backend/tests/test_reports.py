"""PowerPoint report generation: default deck and client template filling."""

from __future__ import annotations

import io

from conftest import run_analysis
from pptx import Presentation
from pptx.util import Inches

PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _all_text(prs) -> str:
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
            if shape.has_table:
                out += [c.text for row in shape.table.rows for c in row.cells]
    return "\n".join(out)


def test_default_deck(client, alice, engagement_with_evidence):
    eid = engagement_with_evidence["id"]
    run_analysis(client, alice, eid, "samm")
    client.put(f"/api/engagements/{eid}/overrides/samm/G-SM", headers=alice,
               json={"score": 1, "reason": "Strategy not yet approved"})
    r = client.get(f"/api/engagements/{eid}/report", headers=alice)
    assert r.status_code == 200
    assert r.headers["content-type"] == PPTX
    assert "Acme-Payments" in r.headers["content-disposition"]
    prs = Presentation(io.BytesIO(r.content))
    text = _all_text(prs)
    for expected in ("Executive summary", "Scope and method", "SAMM maturity by domain", "SAMM practice scores",
                     "SSDF", "CSF 2.0", "Key gaps", "30 / 60 / 90-day roadmap", "Appendix: evidence",
                     "Interview with lead developer", "Payments API and web app", "Override"):
        assert expected in text, expected
    charts = [s for slide in prs.slides for s in slide.shapes if s.has_chart]
    assert len(charts) == 3  # one radar per selected framework
    audit = client.get(f"/api/audit?engagement_id={eid}", headers=alice).json()
    assert audit[0]["action"] == "report.download"


def test_default_deck_without_analysis(client, alice, engagement):
    r = client.get(f"/api/engagements/{engagement['id']}/report", headers=alice)
    assert r.status_code == 200
    Presentation(io.BytesIO(r.content))


def _template() -> bytes:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = s.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1))
    p = box.text_frame.paragraphs[0]
    # token split across runs, as PowerPoint often does
    for piece in ("Report for {{cli", "ent_name}} / {{app_name}} on {{date}}"):
        p.add_run().text = piece
    s.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(8), Inches(1)).text_frame.text = (
        "Docs: {{document_count}}, interviews: {{interview_count}}, overall: {{overall_maturity}}, "
        "frameworks: {{frameworks}}, scope: {{scope}}, unknown: {{nope}}")
    s.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(8), Inches(1)).text_frame.text = "{{executive_summary}}"
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    s2.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(4)).text_frame.text = "{{chart:samm}}"
    s2.shapes.add_textbox(Inches(6.5), Inches(1), Inches(3), Inches(5)).text_frame.text = " {{table:scores:SSDF}} "
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    s3.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(3)).text_frame.text = "{{table:gaps}}"
    s3.shapes.add_textbox(Inches(0.5), Inches(4), Inches(9), Inches(3)).text_frame.text = "{{table:roadmap}}"
    s3.shapes.add_textbox(Inches(0.5), Inches(7), Inches(3), Inches(0.5)).text_frame.text = "{{chart:bsimm}}"
    tbl = s3.shapes.add_table(1, 1, Inches(4), Inches(7), Inches(3), Inches(0.5)).table
    tbl.cell(0, 0).text = "Client: {{client_name}}"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_template_filling(client, alice, engagement_with_evidence):
    eid = engagement_with_evidence["id"]
    run_analysis(client, alice, eid, "samm")
    r = client.post("/api/templates", headers=alice, files={"file": ("client-template.pptx", _template(), PPTX)})
    assert r.status_code == 201, r.text
    tpl = r.json()
    assert {"client_name", "app_name", "chart:samm", "table:scores:SSDF", "table:gaps", "table:roadmap"} <= set(
        tpl["tokens"])
    r = client.get(f"/api/engagements/{eid}/report?template_id={tpl['id']}", headers=alice)
    assert r.status_code == 200
    prs = Presentation(io.BytesIO(r.content))
    text = _all_text(prs)
    assert "Report for Acme / Payments on" in text
    assert "Docs: 1, interviews: 1" in text
    assert "SAMM:" in text and "(projected)" in text
    assert "scope: Payments API and web app" in text
    assert "{{nope}}" in text  # unknown tokens are left as-is
    assert "{{" not in text.replace("{{nope}}", "")
    assert "Client: Acme" in text
    s2 = prs.slides[1]
    charts = [s for s in s2.shapes if s.has_chart]
    tables = [s for s in s2.shapes if s.has_table]
    assert len(charts) == 1 and len(tables) == 1
    assert (charts[0].left, charts[0].top, charts[0].width, charts[0].height) == (
        Inches(1), Inches(1), Inches(5), Inches(4))
    assert (tables[0].left, tables[0].top) == (Inches(6.5), Inches(1))
    assert "PO.1" in "\n".join(c.text for row in tables[0].table.rows for c in row.cells)
    s3 = prs.slides[2]
    headers = [s.table.cell(0, 0).text for s in s3.shapes if s.has_table]
    assert "Framework" in headers and "Horizon" in headers
    assert "not part of this engagement" in _all_text(prs)  # bsimm is not selected


def test_template_upload_validation(client, alice):
    r = client.post("/api/templates", headers=alice, files={"file": ("x.pdf", b"%PDF", "application/pdf")})
    assert r.status_code == 415
    r = client.post("/api/templates", headers=alice, files={"file": ("x.pptx", b"PKnotreally", PPTX)})
    assert r.status_code == 415


def test_sample_template_round_trip(client, alice, engagement_with_evidence):
    sample = client.get("/api/templates/sample", headers=alice)
    assert sample.status_code == 200
    tpl = client.post("/api/templates", headers=alice, files={"file": ("sample.pptx", sample.content, PPTX)}).json()
    assert len(tpl["tokens"]) == 13
    r = client.get(f"/api/engagements/{engagement_with_evidence['id']}/report?template_id={tpl['id']}", headers=alice)
    assert "{{" not in _all_text(Presentation(io.BytesIO(r.content)))
    assert client.delete(f"/api/templates/{tpl['id']}", headers=alice).status_code == 204

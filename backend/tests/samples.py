"""Build small sample files of every supported type in memory."""

from __future__ import annotations

import io

PDF_TEXT = "Threat modeling is performed with STRIDE for every new feature."


def make_pdf(text: str = PDF_TEXT) -> bytes:
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Secure development policy", level=1)
    d.add_paragraph("All code changes require peer review before merge.")
    d.add_heading("Tooling", level=2)
    d.add_paragraph("SAST runs on every pull request.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Tool"
    t.cell(0, 1).text = "Purpose"
    t.cell(1, 0).text = "Semgrep"
    t.cell(1, 1).text = "Static analysis"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "CI/CD security"
    s.placeholders[1].text = "Branch protection is enforced on main."
    s.notes_slide.notes_text_frame.text = "Pipeline runners use OIDC."
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Findings"
    tbl = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    tbl.cell(0, 0).text = "Finding"
    tbl.cell(1, 0).text = "No SBOM generated"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_xlsx() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Controls"
    ws.append(["Control", "Status"])
    ws.append(["Secrets scanning", "Implemented"])
    ws.append(["DAST", "Planned"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


EML = (
    b"From: Security Lead <lead@acme.test>\r\n"
    b"To: team@acme.test\r\n"
    b"Subject: Pen test schedule\r\n"
    b"Date: Tue, 1 Sep 2026 10:00:00 +0000\r\n"
    b"Content-Type: text/plain; charset=utf-8\r\n\r\n"
    b"The annual penetration test is booked for October. Call 020 7946 0958 with questions.\r\n"
)

SAMPLES: dict[str, tuple[bytes, str]] = {
    "policy.pdf": (make_pdf(), "STRIDE"),
    "policy.docx": (make_docx(), "# Secure development policy"),
    "deck.pptx": (make_pptx(), "Branch protection"),
    "controls.xlsx": (make_xlsx(), "Secrets scanning | Implemented"),
    "notes.md": (b"# Notes\n\nWe use Dependabot.", "Dependabot"),
    "notes.txt": (b"Plain text evidence about incident response.", "incident response"),
    "inventory.csv": (b"app,owner\npayments,team-a\n", "payments | team-a"),
    "mail.eml": (EML, "penetration test"),
    "pipeline.yaml": (b"# CI pipeline\nsteps:\n  - run: semgrep --config auto\n", "semgrep"),
    "config.json": (b'{"sast": {"enabled": true}}', '"enabled": true'),
}

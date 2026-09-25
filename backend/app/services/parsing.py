"""Convert uploaded evidence files to heading-annotated plain text.

Headings are emitted as Markdown ``#`` lines so the chunker can split on them.
"""

from __future__ import annotations

import csv
import email
import email.policy
import io
import json
import re
import zipfile
from dataclasses import dataclass
from html import unescape
from pathlib import PurePath

import yaml

MAX_TEXT_CHARS = 2_000_000
MAX_ZIP_UNCOMPRESSED = 250 * 1024 * 1024
MAX_ZIP_ENTRIES = 5000
MAX_SHEET_ROWS = 5000

# extension -> (canonical content type, parser kind)
ALLOWED_TYPES: dict[str, tuple[str, str]] = {
    ".pdf": ("application/pdf", "pdf"),
    ".docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    ".pptx": ("application/vnd.openxmlformats-officedocument.presentationml.presentation", "pptx"),
    ".xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    ".md": ("text/markdown", "text"),
    ".markdown": ("text/markdown", "text"),
    ".txt": ("text/plain", "text"),
    ".csv": ("text/csv", "csv"),
    ".eml": ("message/rfc822", "eml"),
    ".yaml": ("application/yaml", "yaml"),
    ".yml": ("application/yaml", "yaml"),
    ".json": ("application/json", "json"),
}


class UnsupportedFileError(ValueError):
    pass


@dataclass
class ParsedFile:
    text: str
    content_type: str
    kind: str


def classify(filename: str) -> tuple[str, str]:
    ext = PurePath(filename.lower()).suffix
    if ext not in ALLOWED_TYPES:
        allowed = ", ".join(sorted({e.lstrip(".").upper() for e in ALLOWED_TYPES}))
        raise UnsupportedFileError(f"File type '{ext or 'none'}' is not allowed. Allowed: {allowed}")
    return ALLOWED_TYPES[ext]


def safe_filename(name: str) -> str:
    name = PurePath(name.replace("\\", "/")).name
    name = re.sub(r"[^\w.\- ()]+", "_", name).strip() or "upload"
    return name[:200]


def _check_zip(data: bytes) -> None:
    if not data.startswith(b"PK"):
        raise UnsupportedFileError("File content does not match its Office Open XML extension")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
            if len(infos) > MAX_ZIP_ENTRIES:
                raise UnsupportedFileError("Archive has too many entries")
            if sum(i.file_size for i in infos) > MAX_ZIP_UNCOMPRESSED:
                raise UnsupportedFileError("Archive expands beyond the allowed size")
    except zipfile.BadZipFile as exc:
        raise UnsupportedFileError("Corrupt Office document") from exc


def _decode(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _looks_binary(data: bytes) -> bool:
    return b"\x00" in data[:4096] and data[:2] not in (b"\xff\xfe", b"\xfe\xff")


def _parse_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    if not data.lstrip()[:5].startswith(b"%PDF"):
        raise UnsupportedFileError("File content is not a PDF")
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:  # noqa: BLE001
                raise UnsupportedFileError("Encrypted PDFs are not supported") from exc
        parts = []
        for i, page in enumerate(reader.pages, start=1):
            txt = (page.extract_text() or "").strip()
            if txt:
                parts.append(f"## Page {i}\n\n{txt}")
    except UnsupportedFileError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise UnsupportedFileError(f"Could not read PDF: {exc}") from exc
    return "\n\n".join(parts)


def _parse_docx(data: bytes) -> str:
    import docx

    _check_zip(data)
    document = docx.Document(io.BytesIO(data))
    out: list[str] = []
    body = document.element.body
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, document)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            m = re.match(r"Heading (\d)", style)
            if m:
                out.append("#" * min(int(m.group(1)), 6) + " " + text)
            elif style == "Title":
                out.append("# " + text)
            else:
                out.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    out.append(" | ".join(cells))
    return "\n\n".join(out)


def _parse_pptx(data: bytes) -> str:
    from pptx import Presentation

    _check_zip(data)
    prs = Presentation(io.BytesIO(data))
    out: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        out.append(f"## Slide {i}" + (f": {title}" if title else ""))
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt:
                    out.append(txt)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        out.append(" | ".join(cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append(f"Speaker notes: {notes}")
    return "\n\n".join(out)


def _parse_xlsx(data: bytes) -> str:
    import openpyxl

    _check_zip(data)
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: list[str] = []
    try:
        for ws in wb.worksheets:
            out.append(f"## Sheet: {ws.title}")
            rows = []
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n >= MAX_SHEET_ROWS:
                    rows.append(f"(truncated after {MAX_SHEET_ROWS} rows)")
                    break
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    rows.append(" | ".join(cells).rstrip(" |"))
            out.append("\n".join(rows))
    finally:
        wb.close()
    return "\n\n".join(out)


def _parse_csv(data: bytes) -> str:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096])
    except csv.Error:
        dialect = csv.excel
    rows = []
    for n, row in enumerate(csv.reader(io.StringIO(text), dialect)):
        if n >= MAX_SHEET_ROWS:
            break
        if any(c.strip() for c in row):
            rows.append(" | ".join(c.strip() for c in row))
    return "\n".join(rows)


_TAG_RE = re.compile(r"<(script|style)[^>]*>.*?</\1>|<[^>]+>", re.S | re.I)


def _html_to_text(html: str) -> str:
    html = re.sub(r"<br\s*/?>|</p>|</div>|</li>|</h\d>", "\n", html, flags=re.I)
    return re.sub(r"\n{3,}", "\n\n", unescape(_TAG_RE.sub("", html))).strip()


def _parse_eml(data: bytes) -> str:
    msg = email.message_from_bytes(data, policy=email.policy.default)
    header = [f"# Email: {msg.get('subject', '(no subject)')}"]
    for h in ("from", "to", "cc", "date"):
        if msg.get(h):
            header.append(f"{h.title()}: {msg.get(h)}")
    body = ""
    part = msg.get_body(preferencelist=("plain", "html"))
    if part is not None:
        content = part.get_content()
        body = _html_to_text(content) if part.get_content_type() == "text/html" else content
    attachments = [a.get_filename() for a in msg.iter_attachments() if a.get_filename()]
    if attachments:
        body += "\n\nAttachments: " + ", ".join(attachments)
    return "\n".join(header) + "\n\n" + body.strip()


def _parse_json(data: bytes) -> str:
    text = _decode(data)
    try:
        return json.dumps(json.loads(text), indent=2, ensure_ascii=False)
    except json.JSONDecodeError:
        return text


def _parse_yaml(data: bytes) -> str:
    text = _decode(data)
    try:
        list(yaml.safe_load_all(text))  # validate, but keep original (comments carry meaning)
    except yaml.YAMLError:
        pass
    return text


def parse_file(filename: str, data: bytes) -> ParsedFile:
    content_type, kind = classify(filename)
    if kind in {"text", "csv", "json", "yaml", "eml"} and _looks_binary(data):
        raise UnsupportedFileError("File looks binary but has a text extension")
    parser = {
        "pdf": _parse_pdf,
        "docx": _parse_docx,
        "pptx": _parse_pptx,
        "xlsx": _parse_xlsx,
        "csv": _parse_csv,
        "eml": _parse_eml,
        "json": _parse_json,
        "yaml": _parse_yaml,
        "text": _decode,
    }[kind]
    try:
        text = parser(data)
    except UnsupportedFileError:
        raise
    except Exception as exc:  # noqa: BLE001 - surface a clean error to the user
        raise UnsupportedFileError(f"Could not parse {kind.upper()} file: {exc}") from exc
    text = text.replace("\r\n", "\n").replace("\x00", "")
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS]
    return ParsedFile(text=text.strip(), content_type=content_type, kind=kind)

"""PowerPoint reports: a built-in Deloitte-branded deck and client template filling."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# Deloitte palette
GREEN = RGBColor(0x86, 0xBC, 0x25)
GREEN_6 = RGBColor(0x26, 0x89, 0x0D)
GREEN_7 = RGBColor(0x04, 0x6A, 0x38)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x53, 0x56, 0x5A)
LIGHT_GREY = RGBColor(0xD0, 0xD0, 0xCE)
PALE = RGBColor(0xF3, 0xF6, 0xEC)
FONT = "Arial"

TEXT_TOKENS = (
    "client_name",
    "app_name",
    "scope",
    "date",
    "executive_summary",
    "frameworks",
    "overall_maturity",
    "document_count",
    "interview_count",
)
TOKEN_RE = re.compile(r"\{\{\s*([a-z_]+(?::[A-Za-z0-9_.\- ]+)*)\s*\}\}")


@dataclass
class ReportContext:
    client_name: str
    app_name: str
    business_unit: str | None
    scope: str
    status: str
    report_date: date
    results: list[dict[str, Any]]  # build_results() output per selected framework
    gaps: list[dict[str, Any]]
    roadmap: list[dict[str, Any]]
    documents: list[dict[str, Any]]
    analyzers: list[str]
    override_count: int
    executive_summary: str = ""
    framework_aliases: dict[str, str] = field(default_factory=dict)

    @property
    def document_count(self) -> int:
        return sum(1 for d in self.documents if d["kind"] == "file")

    @property
    def interview_count(self) -> int:
        return sum(1 for d in self.documents if d["kind"] == "interview")

    def result_for(self, name: str) -> dict[str, Any] | None:
        key = self.framework_aliases.get(_alias(name))
        return next((r for r in self.results if r["framework"]["key"] == key), None)

    def overall_text(self) -> str:
        parts = []
        for r in self.results:
            fw = r["framework"]
            val = "n/a" if r["overall"] is None else f"{r['overall']:.1f} / {fw['scale']['max']:g}"
            parts.append(f"{fw['short_name']}: {val}" + (" (projected)" if r["is_projected"] else ""))
        return "; ".join(parts) if parts else "n/a"


def _alias(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", s.lower())


def build_aliases(frameworks: list[Any]) -> dict[str, str]:
    """Map loose names ('samm', 'OWASP SAMM', 'nist_ssdf', 'SSDF') to framework keys."""
    aliases: dict[str, str] = {}
    for fw in frameworks:
        for name in (fw.key, fw.name, fw.short_name, fw.short_name.split(" ")[0], fw.key.split("_")[-1]):
            aliases.setdefault(_alias(name), fw.key)
    return aliases


def make_executive_summary(ctx: ReportContext) -> str:
    if not ctx.results:
        return "No frameworks were selected for this engagement."
    lines = []
    analyzed = [r for r in ctx.results if not r["is_projected"]]
    lead = analyzed[0] if analyzed else ctx.results[0]
    fw = lead["framework"]
    if lead["overall"] is not None:
        lines.append(
            f"{ctx.app_name} ({ctx.client_name}) achieves an overall {fw['short_name']} maturity of "
            f"{lead['overall']:.1f} out of {fw['scale']['max']:g} against a target of {lead['target']:g}."
        )
    doms = [d for d in lead["domains"] if d["current"] is not None]
    if doms:
        best = max(doms, key=lambda d: d["current"])
        worst = min(doms, key=lambda d: d["current"])
        lines.append(
            f"The strongest area is {best['name']} ({best['current']:.1f}); the weakest is "
            f"{worst['name']} ({worst['current']:.1f})."
        )
    others = [r for r in ctx.results if r is not lead and r["overall"] is not None]
    if others:
        lines.append(
            "Other frameworks: "
            + "; ".join(
                f"{r['framework']['short_name']} {r['overall']:.1f}/{r['framework']['scale']['max']:g}"
                + (" (projected)" if r["is_projected"] else "")
                for r in others
            )
            + "."
        )
    if ctx.gaps:
        top = ", ".join(g["practice"] for g in ctx.gaps[:3])
        lines.append(f"Priority gaps: {top}.")
    if ctx.roadmap:
        n30 = sum(1 for r in ctx.roadmap if r["horizon"] == 30)
        lines.append(f"The roadmap contains {len(ctx.roadmap)} actions, {n30} of them within 30 days.")
    lines.append(
        f"Evidence: {ctx.document_count} document(s) and {ctx.interview_count} interview(s); "
        f"{ctx.override_count} score(s) adjusted by assessors."
    )
    return " ".join(lines)


def text_values(ctx: ReportContext) -> dict[str, str]:
    return {
        "client_name": ctx.client_name,
        "app_name": ctx.app_name,
        "scope": ctx.scope or "-",
        "date": ctx.report_date.strftime("%d %B %Y"),
        "executive_summary": ctx.executive_summary,
        "frameworks": ", ".join(r["framework"]["name"] + " " + r["framework"]["version"] for r in ctx.results),
        "overall_maturity": ctx.overall_text(),
        "document_count": str(ctx.document_count),
        "interview_count": str(ctx.interview_count),
    }


# --- drawing primitives ------------------------------------------------------------------


def _style_run(run, size: float, color: RGBColor = BLACK, bold: bool = False) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text(slide, left, top, width, height, text, size=14, color=BLACK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        _style_run(p.add_run(), size, color, bold)
        p.runs[0].text = line
    return box


def add_radar_chart(slide, left, top, width, height, result: dict[str, Any]):
    domains = result["domains"]
    data = CategoryChartData()
    data.categories = [d["name"] for d in domains]
    data.add_series("Current", [d["current"] or 0 for d in domains])
    data.add_series("Target", [d["target"] for d in domains])
    chart_type = XL_CHART_TYPE.RADAR_MARKERS if len(domains) >= 3 else XL_CHART_TYPE.BAR_CLUSTERED
    gframe = slide.shapes.add_chart(chart_type, left, top, width, height, data)
    chart = gframe.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(10)
    chart.legend.font.name = FONT
    chart.font.name = FONT
    chart.font.size = Pt(9)
    for series, color in zip(chart.plots[0].series, (GREEN, BLACK), strict=True):
        fmt = series.format
        fmt.line.color.rgb = color
        fmt.line.width = Pt(2.25)
        if chart_type == XL_CHART_TYPE.BAR_CLUSTERED:
            fmt.fill.solid()
            fmt.fill.fore_color.rgb = color
        else:
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = color
            series.marker.format.line.color.rgb = color
    va = chart.value_axis
    va.minimum_scale = result["framework"]["scale"]["min"]
    va.maximum_scale = result["framework"]["scale"]["max"]
    va.major_unit = 1
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = LIGHT_GREY
    va.tick_labels.font.size = Pt(8)
    return gframe


def _set_cell(cell, text, size, bold=False, color=BLACK, fill: RGBColor | None = None, align=PP_ALIGN.LEFT):
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text)
    _style_run(run, size, color, bold)
    cell.margin_left = cell.margin_right = Inches(0.05)
    cell.margin_top = cell.margin_bottom = Inches(0.02)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    else:
        cell.fill.background()


HEAT_RAMP = [  # sequential single-hue ramp (light -> dark), matches the web heatmap
    (0.0, RGBColor(0xF1, 0xF6, 0xE6), BLACK),
    (0.2, RGBColor(0xDC, 0xEB, 0xC0), BLACK),
    (0.4, RGBColor(0xB5, 0xD7, 0x7A), BLACK),
    (0.6, GREEN, BLACK),
    (0.8, GREEN_6, WHITE),
    (0.95, GREEN_7, WHITE),
]


def heat_color(score: float | None, smin: float, smax: float) -> tuple[RGBColor, RGBColor]:
    """(fill, text) colours for a score cell."""
    if score is None:
        return RGBColor(0xED, 0xED, 0xF0), GREY
    n = (score - smin) / (smax - smin) if smax > smin else 0
    fill, text = HEAT_RAMP[0][1], HEAT_RAMP[0][2]
    for threshold, f, t in HEAT_RAMP:
        if n >= threshold - 1e-9:
            fill, text = f, t
    return fill, text


def add_table(slide, left, top, width, height, headers: list[str], rows: list[list[Any]],
              col_widths: list[float] | None = None, font_size: float | None = None,
              score_cols: dict[int, tuple[float, float]] | None = None):
    n_rows = len(rows) + 1
    size = font_size or max(7.0, min(11.0, (Emu(height).pt / n_rows) * 0.45))
    gframe = slide.shapes.add_table(n_rows, len(headers), left, top, width, height)
    table = gframe.table
    tbl_pr = gframe._element.graphic.graphicData.tbl.tblPr
    tbl_pr.set("bandRow", "0")
    tbl_pr.set("firstRow", "1")
    total = sum(col_widths) if col_widths else len(headers)
    for i in range(len(headers)):
        frac = (col_widths[i] if col_widths else 1) / total
        table.columns[i].width = int(Emu(width) * frac)
    row_h = int(Emu(height) / n_rows)
    for r in range(n_rows):
        table.rows[r].height = row_h
    for c, h in enumerate(headers):
        _set_cell(table.cell(0, c), h, size, bold=True, color=WHITE, fill=BLACK)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            fill = PALE if r % 2 == 0 else WHITE
            color = BLACK
            if score_cols and c in score_cols and isinstance(value, (int, float)):
                fill, color = heat_color(value, *score_cols[c])
                value = f"{value:.1f}"
            elif value is None:
                value = "-"
            _set_cell(table.cell(r, c), value, size, fill=fill, color=color,
                      align=PP_ALIGN.CENTER if score_cols and c in score_cols else PP_ALIGN.LEFT)
    return gframe


def score_rows(result: dict[str, Any]) -> list[list[Any]]:
    dom_names = {d["id"]: d["name"] for d in result["domains"]}
    status_label = {"assessed": "Assessed", "projected": "Projected", "overridden": "Override",
                    "not_assessed": "Not assessed"}
    return [
        [dom_names[p["domain_id"]], f"{p['id']} {p['name']}", p["score"], p["target"], status_label[p["status"]]]
        for p in result["practices"]
    ]


SCORE_HEADERS = ["Domain", "Practice", "Score", "Target", "Status"]
SCORE_WIDTHS = [2.2, 5.0, 0.9, 0.9, 1.3]
GAP_HEADERS = ["Framework", "Practice", "Score", "Target", "Gap"]
GAP_WIDTHS = [1.2, 3.2, 0.8, 0.8, 5.0]
ROADMAP_HEADERS = ["Horizon", "Priority", "Action", "Practices"]
ROADMAP_WIDTHS = [0.9, 0.9, 6.2, 2.4]


def gap_rows(ctx: ReportContext, limit: int = 10) -> list[list[Any]]:
    return [[g["framework"], g["practice"], f"{g['score']:.1f}", f"{g['target']:g}", g["gap"]]
            for g in ctx.gaps[:limit]]


def roadmap_rows(ctx: ReportContext, limit: int = 12) -> list[list[Any]]:
    return [
        [f"{r['horizon']} days", r["priority"].title(), r["text"], ", ".join(r["practices"][:3])]
        for r in ctx.roadmap[:limit]
    ]


# --- default deck ------------------------------------------------------------------------

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _header(slide, title: str, kicker: str = "") -> None:
    if kicker:
        add_text(slide, MARGIN, Inches(0.35), W - 2 * MARGIN, Inches(0.3), kicker.upper(), 10, GREEN_6, True)
    box = add_text(slide, MARGIN, Inches(0.6), W - 2 * MARGIN, Inches(0.8), title, 28, BLACK, True)
    p = box.text_frame.paragraphs[0]
    dot = p.add_run()
    dot.text = "."
    _style_run(dot, 28, GREEN, True)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, Inches(1.4), Inches(0.8), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREEN
    bar.line.fill.background()


def _footer(slide, ctx: ReportContext, n: int) -> None:
    add_text(slide, MARGIN, H - Inches(0.45), Inches(8), Inches(0.3),
             f"{ctx.client_name} | {ctx.app_name} | SSDLC maturity assessment | Confidential", 8, GREY)
    add_text(slide, W - MARGIN - Inches(1), H - Inches(0.45), Inches(1), Inches(0.3), str(n), 8, GREY,
             align=PP_ALIGN.RIGHT)


def _kpi(slide, left, top, value: str, label: str) -> None:
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.8), Inches(1.3))
    card.adjustments[0] = 0.12
    card.fill.solid()
    card.fill.fore_color.rgb = PALE
    card.line.fill.background()
    add_text(slide, left + Inches(0.2), top + Inches(0.12), Inches(2.4), Inches(0.6), value, 26, GREEN_7, True)
    add_text(slide, left + Inches(0.2), top + Inches(0.78), Inches(2.4), Inches(0.4), label, 11, GREY)


def build_default_deck(ctx: ReportContext) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    n = 0

    # 1. Title
    s = _blank(prs)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = BLACK
    add_text(s, MARGIN, Inches(1.0), Inches(8), Inches(0.4), "SECURE SDLC MATURITY ASSESSMENT", 12, GREEN, True)
    title = add_text(s, MARGIN, Inches(2.2), W - 2 * MARGIN, Inches(1.6), ctx.app_name, 48, WHITE, True)
    dot = title.text_frame.paragraphs[0].add_run()
    dot.text = "."
    _style_run(dot, 48, GREEN, True)
    add_text(s, MARGIN, Inches(3.9), W - 2 * MARGIN, Inches(0.6),
             ctx.client_name + (f" | {ctx.business_unit}" if ctx.business_unit else ""), 22, WHITE)
    add_text(s, MARGIN, Inches(5.8), Inches(8), Inches(0.4),
             ctx.report_date.strftime("%d %B %Y") + f" | Status: {ctx.status.replace('_', ' ').title()}", 12,
             LIGHT_GREY)

    # 2. Executive summary
    n += 1
    s = _blank(prs)
    _header(s, "Executive summary", "Overview")
    lead = ctx.results[0] if ctx.results else None
    _kpi(s, MARGIN, Inches(1.8),
         f"{lead['overall']:.1f}" if lead and lead["overall"] is not None else "n/a",
         f"Overall {lead['framework']['short_name']} maturity" if lead else "Overall maturity")
    _kpi(s, MARGIN + Inches(3.05), Inches(1.8), str(len(ctx.results)), "Frameworks assessed")
    _kpi(s, MARGIN + Inches(6.1), Inches(1.8), str(ctx.document_count), "Documents reviewed")
    _kpi(s, MARGIN + Inches(9.15), Inches(1.8), str(ctx.interview_count), "Interviews held")
    add_text(s, MARGIN, Inches(3.4), W - 2 * MARGIN, Inches(3.4), ctx.executive_summary, 14, BLACK)
    _footer(s, ctx, n)

    # 3. Scope and method
    n += 1
    s = _blank(prs)
    _header(s, "Scope and method", "Approach")
    add_text(s, MARGIN, Inches(1.8), Inches(5.6), Inches(0.4), "Scope", 16, GREEN_7, True)
    add_text(s, MARGIN, Inches(2.25), Inches(5.6), Inches(4.2), ctx.scope or "Not specified.", 12)
    analyzer_text = ", ".join(sorted(set(ctx.analyzers))) or "none yet"
    method = [
        f"{ctx.document_count} client document(s) and {ctx.interview_count} interview note(s) were ingested, "
        "redacted (e-mail addresses and phone numbers) and split into heading-aware evidence chunks.",
        f"Practices were scored per framework domain by: {analyzer_text}. Every citation was verified to "
        "appear word for word in the source evidence.",
        "Frameworks not analysed directly are projected through a shared taxonomy of SSDLC capabilities.",
        f"Assessors reviewed the results; {ctx.override_count} score(s) were overridden with a documented reason.",
        "Frameworks: " + ", ".join(r["framework"]["name"] + " " + r["framework"]["version"] for r in ctx.results) + ".",
    ]
    add_text(s, Inches(6.9), Inches(1.8), Inches(5.8), Inches(0.4), "Method", 16, GREEN_7, True)
    add_text(s, Inches(6.9), Inches(2.25), Inches(5.8), Inches(4.2), "\n\n".join("- " + m for m in method), 12)
    _footer(s, ctx, n)

    # 4. Per framework: chart + score tables
    for r in ctx.results:
        fw = r["framework"]
        n += 1
        s = _blank(prs)
        _header(s, f"{fw['short_name']} maturity by domain", fw["name"] + (" - projected" if r["is_projected"] else ""))
        add_radar_chart(s, MARGIN, Inches(1.7), Inches(7.2), Inches(5.2), r)
        rows = [[d["name"], d["current"], d["target"]] for d in r["domains"]]
        add_table(s, Inches(8.1), Inches(1.9), Inches(4.6), Inches(0.4) * (len(rows) + 1),
                  ["Domain", "Current", "Target"], rows, [2.6, 1, 1], 10,
                  score_cols={1: (fw["scale"]["min"], fw["scale"]["max"])})
        note = ("Scores projected from other frameworks via the capability taxonomy."
                if r["is_projected"] else f"Overall: {r['overall'] if r['overall'] is not None else 'n/a'} "
                f"on a {fw['scale']['min']:g}-{fw['scale']['max']:g} scale.")
        add_text(s, Inches(8.1), Inches(6.2), Inches(4.6), Inches(0.6), note, 10, GREY)
        _footer(s, ctx, n)

        rows = score_rows(r)
        per_page = 11
        for page in range(0, len(rows), per_page):
            n += 1
            s = _blank(prs)
            suffix = f" ({page // per_page + 1}/{(len(rows) - 1) // per_page + 1})" if len(rows) > per_page else ""
            _header(s, f"{fw['short_name']} practice scores{suffix}", fw["name"])
            chunk = rows[page : page + per_page]
            add_table(s, MARGIN, Inches(1.7), W - 2 * MARGIN, Inches(0.4) * (len(chunk) + 1), SCORE_HEADERS,
                      chunk, SCORE_WIDTHS, 10, score_cols={2: (fw["scale"]["min"], fw["scale"]["max"])})
            _footer(s, ctx, n)

    # 5. Key gaps
    n += 1
    s = _blank(prs)
    _header(s, "Key gaps", "Findings")
    rows = gap_rows(ctx, 10)
    if rows:
        add_table(s, MARGIN, Inches(1.7), W - 2 * MARGIN, Inches(0.46) * (len(rows) + 1), GAP_HEADERS, rows,
                  GAP_WIDTHS, 9)
    else:
        add_text(s, MARGIN, Inches(1.8), Inches(8), Inches(0.5), "No practices fall below target.", 14)
    _footer(s, ctx, n)

    # 6. Roadmap
    n += 1
    s = _blank(prs)
    _header(s, "30 / 60 / 90-day roadmap", "Recommendations")
    col_w = (W - 2 * MARGIN - Inches(0.5)) / 3
    for i, horizon in enumerate(HORIZON_COLUMNS):
        left = MARGIN + i * (col_w + Inches(0.25))
        head = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.7), col_w, Inches(0.55))
        head.fill.solid()
        head.fill.fore_color.rgb = (GREEN, GREEN_6, GREEN_7)[i]
        head.line.fill.background()
        tf = head.text_frame
        tf.text = f"Next {horizon} days"
        _style_run(tf.paragraphs[0].runs[0], 14, WHITE, True)
        all_items = [r for r in ctx.roadmap if r["horizon"] == horizon]
        items = all_items[:6]
        text = "\n".join(f"[{r['priority'].title()}] {r['text']}" for r in items) or "No actions."
        if len(all_items) > len(items):
            text += f"\n+ {len(all_items) - len(items)} more action(s) - see the platform for the full roadmap."
        box = add_text(s, left, Inches(2.4), col_w, Inches(4.4), text, 10)
        for p in box.text_frame.paragraphs:
            p.space_after = Pt(6)
    _footer(s, ctx, n)

    # 7. Evidence appendix
    n += 1
    s = _blank(prs)
    _header(s, "Appendix: evidence", "Appendix")
    rows = [
        [d["title"], "Interview" if d["kind"] == "interview" else (d.get("filename") or "").rsplit(".", 1)[-1].upper(),
         d.get("interviewee_role") or "-", d.get("date") or "-", d["chunks"]]
        for d in ctx.documents[:14]
    ]
    if rows:
        add_table(s, MARGIN, Inches(1.7), W - 2 * MARGIN, Inches(0.36) * (len(rows) + 1),
                  ["Title", "Type", "Interviewee role", "Date", "Chunks"], rows, [5, 1.2, 2.5, 1.5, 0.9], 9)
        if len(ctx.documents) > 14:
            add_text(s, MARGIN, H - Inches(0.9), Inches(8), Inches(0.3),
                     f"...and {len(ctx.documents) - 14} more item(s).", 9, GREY)
    else:
        add_text(s, MARGIN, Inches(1.8), Inches(8), Inches(0.5), "No evidence uploaded.", 14)
    _footer(s, ctx, n)

    cited = [(r["framework"]["short_name"], p, c) for r in ctx.results for p in r["practices"] for c in p["citations"][:1]]
    if cited:
        n += 1
        s = _blank(prs)
        _header(s, "Appendix: cited evidence", "Appendix")
        rows = [[f"{fw} {p['id']}", f"“{c['quote'][:220]}”", c["document_title"]] for fw, p, c in cited[:10]]
        add_table(s, MARGIN, Inches(1.7), W - 2 * MARGIN, Inches(0.46) * (len(rows) + 1),
                  ["Practice", "Quote", "Source"], rows, [1.6, 7.4, 2.6], 9)
        _footer(s, ctx, n)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


HORIZON_COLUMNS = (30, 60, 90)


# --- template filling --------------------------------------------------------------------


def _iter_shapes(shapes):
    for shape in list(shapes):
        yield shape
        if shape.shape_type is not None and hasattr(shape, "shapes"):  # group shape
            yield from _iter_shapes(shape.shapes)


def _replace_in_paragraph(paragraph, values: dict[str, str]) -> None:
    runs = paragraph.runs
    if not runs:
        return
    full = "".join(r.text for r in runs)
    if "{{" not in full:
        return

    def sub(m: re.Match[str]) -> str:
        token = m.group(1)
        return values.get(token, m.group(0))

    new = TOKEN_RE.sub(sub, full)
    if new == full:
        return
    runs[0].text = new
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def _replace_text_frame(tf, values: dict[str, str]) -> None:
    for p in tf.paragraphs:
        _replace_in_paragraph(p, values)


def find_tokens(data: bytes) -> list[str]:
    prs = Presentation(io.BytesIO(data))
    found: list[str] = []
    for slide in prs.slides:
        for shape in _iter_shapes(slide.shapes):
            texts = []
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                texts += [c.text for row in shape.table.rows for c in row.cells]
            for t in texts:
                for m in TOKEN_RE.finditer(t):
                    if m.group(1) not in found:
                        found.append(m.group(1))
    return found


def _block_token(shape) -> str | None:
    if not shape.has_text_frame:
        return None
    m = TOKEN_RE.fullmatch(shape.text_frame.text.strip())
    if m and m.group(1).split(":")[0] in ("chart", "table"):
        return m.group(1)
    return None


def _render_block(slide, token: str, left, top, width, height, ctx: ReportContext) -> None:
    parts = token.split(":")
    kind = parts[0]
    if kind == "chart" and len(parts) >= 2:
        result = ctx.result_for(":".join(parts[1:]))
        if result is None:
            add_text(slide, left, top, width, height, f"Framework '{parts[1]}' is not part of this engagement.", 10, GREY)
            return
        add_radar_chart(slide, left, top, width, height, result)
        return
    if kind == "table" and len(parts) >= 3 and parts[1] == "scores":
        result = ctx.result_for(":".join(parts[2:]))
        if result is None:
            add_text(slide, left, top, width, height, f"Framework '{parts[2]}' is not part of this engagement.", 10, GREY)
            return
        fw = result["framework"]
        add_table(slide, left, top, width, height, SCORE_HEADERS, score_rows(result), SCORE_WIDTHS,
                  score_cols={2: (fw["scale"]["min"], fw["scale"]["max"])})
        return
    if kind == "table" and len(parts) == 2 and parts[1] == "gaps":
        rows = gap_rows(ctx) or [["-", "No practices below target", None, None, "-"]]
        add_table(slide, left, top, width, height, GAP_HEADERS, rows, GAP_WIDTHS)
        return
    if kind == "table" and len(parts) == 2 and parts[1] == "roadmap":
        rows = roadmap_rows(ctx) or [["-", "-", "No actions", "-"]]
        add_table(slide, left, top, width, height, ROADMAP_HEADERS, rows, ROADMAP_WIDTHS)
        return
    add_text(slide, left, top, width, height, f"Unknown token {{{{{token}}}}}", 10, GREY)


def fill_template(data: bytes, ctx: ReportContext) -> bytes:
    prs = Presentation(io.BytesIO(data))
    values = text_values(ctx)
    for slide in prs.slides:
        blocks = []
        for shape in _iter_shapes(slide.shapes):
            token = _block_token(shape)
            if token:
                blocks.append((shape, token))
                continue
            if shape.has_text_frame:
                _replace_text_frame(shape.text_frame, values)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        _replace_text_frame(cell.text_frame, values)
        for shape, token in blocks:
            left, top, width, height = shape.left, shape.top, shape.width, shape.height
            el = shape._element
            el.getparent().remove(el)
            _render_block(slide, token, left, top, width, height, ctx)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_sample_template() -> bytes:
    """A small example template demonstrating every supported token."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    s = _blank(prs)
    add_text(s, MARGIN, Inches(1.5), W - 2 * MARGIN, Inches(1), "{{app_name}} - SSDLC assessment", 36, BLACK, True)
    add_text(s, MARGIN, Inches(2.6), W - 2 * MARGIN, Inches(0.6), "Prepared for {{client_name}} on {{date}}", 18, GREY)
    add_text(s, MARGIN, Inches(3.4), W - 2 * MARGIN, Inches(2.5),
             "Scope: {{scope}}\nFrameworks: {{frameworks}}\nOverall maturity: {{overall_maturity}}\n"
             "Evidence: {{document_count}} documents, {{interview_count}} interviews", 14)
    s = _blank(prs)
    add_text(s, MARGIN, Inches(0.5), W - 2 * MARGIN, Inches(0.8), "Executive summary", 28, BLACK, True)
    add_text(s, MARGIN, Inches(1.5), W - 2 * MARGIN, Inches(4.5), "{{executive_summary}}", 14)
    s = _blank(prs)
    add_text(s, MARGIN, Inches(0.5), W - 2 * MARGIN, Inches(0.8), "SAMM maturity", 28, BLACK, True)
    add_text(s, MARGIN, Inches(1.5), Inches(6), Inches(5), "{{chart:samm}}", 12)
    add_text(s, Inches(7), Inches(1.5), Inches(5.7), Inches(5), "{{table:scores:samm}}", 12)
    s = _blank(prs)
    add_text(s, MARGIN, Inches(0.5), W - 2 * MARGIN, Inches(0.8), "Key gaps", 28, BLACK, True)
    add_text(s, MARGIN, Inches(1.5), W - 2 * MARGIN, Inches(5), "{{table:gaps}}", 12)
    s = _blank(prs)
    add_text(s, MARGIN, Inches(0.5), W - 2 * MARGIN, Inches(0.8), "Roadmap", 28, BLACK, True)
    add_text(s, MARGIN, Inches(1.5), W - 2 * MARGIN, Inches(5), "{{table:roadmap}}", 12)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()

